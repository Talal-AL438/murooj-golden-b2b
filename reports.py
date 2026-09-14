from io import BytesIO
from datetime import datetime
import secrets
from flask import request, render_template, send_file, redirect, url_for, flash, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash, check_password_hash


def register_reports(app, admin_required, db, current_user):
    permission_keys=('reports','agencies','hotels','offers')
    def ensure_permissions_table(c):
        c.execute("CREATE TABLE IF NOT EXISTS staff_permissions (user_id INTEGER PRIMARY KEY,reports INTEGER DEFAULT 0,agencies INTEGER DEFAULT 0,hotels INTEGER DEFAULT 0,offers INTEGER DEFAULT 0,updated_at TEXT,FOREIGN KEY(user_id) REFERENCES users(id))")
    def get_permissions(user_id):
        c=db();ensure_permissions_table(c);row=c.execute("SELECT * FROM staff_permissions WHERE user_id=?",(user_id,)).fetchone();c.close();return {k:bool(row[k]) if row else False for k in permission_keys}
    def audit(c,user_id,action,details=''):
        c.execute("INSERT INTO audit(user_id,action,details,created_at) VALUES(?,?,?,?)",(user_id,action,details,datetime.utcnow().isoformat()))
    def msg(lang,key):
        messages={
            'ar':{'profile':'تم تحديث بيانات الحساب بنجاح.','profile_bad':'يرجى إدخال اسم المسؤول والدولة ورقم WhatsApp بشكل صحيح.','password':'تم تغيير كلمة المرور بنجاح.','bad':'كلمة المرور الحالية غير صحيحة.','short':'يجب أن تكون كلمة المرور الجديدة 8 أحرف على الأقل.','mismatch':'كلمتا المرور الجديدتان غير متطابقتين.','same':'اختر كلمة مرور جديدة مختلفة عن كلمة المرور الحالية.','request_bad':'تعذر إرسال الطلب. تحقق من الفندق والتواريخ وعدد الغرف والأشخاص والبيانات المطلوبة.'},
            'en':{'profile':'Account details updated successfully.','profile_bad':'Please enter a valid contact name, country, and WhatsApp number.','password':'Password changed successfully.','bad':'Current password is incorrect.','short':'New password must be at least 8 characters.','mismatch':'New passwords do not match.','same':'Choose a new password different from your current password.','request_bad':'The request could not be sent. Check the hotel, dates, rooms, persons, and required details.'},
            'id':{'profile':'Data akun berhasil diperbarui.','profile_bad':'Masukkan nama kontak, negara, dan nomor WhatsApp yang valid.','password':'Kata sandi berhasil diubah.','bad':'Kata sandi saat ini salah.','short':'Kata sandi baru minimal 8 karakter.','mismatch':'Konfirmasi kata sandi tidak cocok.','same':'Pilih kata sandi baru yang berbeda dari kata sandi saat ini.','request_bad':'Permintaan tidak dapat dikirim. Periksa hotel, tanggal, kamar, jumlah orang, dan data wajib.'},
            'ms':{'profile':'Maklumat akaun berjaya dikemas kini.','profile_bad':'Masukkan nama pegawai, negara dan nombor WhatsApp yang sah.','password':'Kata laluan berjaya dikemas kini.','bad':'Kata laluan semasa tidak betul.','short':'Kata laluan baharu mestilah sekurang-kurangnya 8 aksara.','mismatch':'Pengesahan kata laluan tidak sepadan.','same':'Pilih kata laluan baharu yang berbeza daripada kata laluan semasa.','request_bad':'Permintaan tidak dapat dihantar. Semak hotel, tarikh, bilik, bilangan orang dan maklumat wajib.'}}
        return messages.get(lang or 'ar',messages['ar'])[key]

    @app.before_request
    def portal_extensions_and_staff_access():
        if '_csrf_token' not in session:session['_csrf_token']=secrets.token_urlsafe(32)
        if request.method=='POST':
            sent=request.form.get('_csrf_token','') or request.headers.get('X-CSRF-Token','')
            if not sent or not secrets.compare_digest(sent,session.get('_csrf_token','')):return 'Invalid or missing CSRF token.',400
        c=db();ensure_permissions_table(c);c.commit();c.close()
        u=current_user()
        if u and not u['active']:
            lang=u['language'] or 'ar';session.clear();flash('تم إيقاف هذا الحساب. تواصل مع الإدارة إذا كنت تحتاج إلى استعادة الوصول.' if lang=='ar' else 'This account has been suspended. Contact the administrator if you need access restored.');return redirect(url_for('login'))
        if request.path=='/request' and request.method=='POST':
            u=current_user()
            if not u or u['role']!='agency':return None
            lang=u['language'] or 'ar';c=db();agency=c.execute("SELECT id FROM agencies WHERE user_id=?",(u['id'],)).fetchone()
            if not agency:c.close();return None
            hv=request.form.get('hotel_id','').strip();city=request.form.get('city','').strip();checkin=request.form.get('checkin','').strip();checkout=request.form.get('checkout','').strip();nationality=request.form.get('nationality','').strip();meal=request.form.get('meal','').strip();notes=request.form.get('notes','').strip();valid=True;hotel_id=None;any_hotel=(hv=='any')
            try:rooms=int(request.form.get('rooms','0'));persons=int(request.form.get('persons','0'))
            except (TypeError,ValueError):rooms=0;persons=0;valid=False
            if rooms<1 or rooms>500 or persons<1 or persons>5000:valid=False
            if city not in ('Makkah','Madinah'):valid=False
            try:
                ci=datetime.strptime(checkin,'%Y-%m-%d').date();co=datetime.strptime(checkout,'%Y-%m-%d').date()
                if co<=ci:valid=False
            except ValueError:valid=False
            if not nationality or len(nationality)>100 or len(notes)>2000:valid=False
            if meal not in ('RO','F.B Indo','F.B Malaysian'):valid=False
            if not any_hotel:
                try:hotel_id=int(hv)
                except (TypeError,ValueError):valid=False
                if hotel_id:
                    hotel=c.execute("SELECT id,city,active FROM hotels WHERE id=?",(hotel_id,)).fetchone()
                    if not hotel or not hotel['active'] or hotel['city']!=city:valid=False
            if not valid:c.close();flash(msg(lang,'request_bad'));return redirect(url_for('new_request'))
            cur=c.execute("INSERT INTO requests(agency_id,hotel_id,any_hotel,city,checkin,checkout,rooms,persons,nationality,meal,notes,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",(agency['id'],hotel_id,1 if any_hotel else 0,city,checkin,checkout,rooms,persons,nationality,meal,notes,'sent',datetime.utcnow().isoformat()));rid=cur.lastrowid;c.execute("INSERT INTO notifications(user_id,type,ref_id,title,body,link,created_at) VALUES(?,?,?,?,?,?,?)",(u['id'],'request',rid,'تحديث الطلب' if lang=='ar' else 'Request update',f'#{rid}','/account',datetime.utcnow().isoformat()));audit(c,u['id'],'request_submit',f"request={rid}, agency={agency['id']}");c.commit();c.close();flash({'ar':'تم إرسال طلبك إلى مروج الذهبية بنجاح. سيتواصل معك فريق الحجوزات عبر WhatsApp لتأكيد التوفر والسعر.','en':'Your request was sent to Murooj Golden successfully. Our reservations team will contact you on WhatsApp to confirm availability and price.','id':'Permintaan Anda berhasil dikirim ke Murooj Golden. Tim reservasi akan menghubungi Anda melalui WhatsApp untuk mengonfirmasi ketersediaan dan harga.','ms':'Permintaan anda berjaya dihantar kepada Murooj Golden. Pasukan tempahan akan menghubungi anda melalui WhatsApp untuk mengesahkan ketersediaan dan harga.'}.get(lang,'Request sent.'));return redirect(url_for('account'))
        if request.path=='/account' and request.method=='POST':
            u=current_user()
            if not u or u['role']!='agency':return None
            lang=u['language'] or 'ar';action=request.form.get('action','');c=db();agency=c.execute("SELECT * FROM agencies WHERE user_id=?",(u['id'],)).fetchone()
            if not agency:c.close();return None
            if action=='profile':
                cn=request.form.get('contact_name','').strip();country=request.form.get('country','').strip();wa=request.form.get('whatsapp','').strip();normalized=''.join(ch for ch in wa if ch.isdigit())
                if len(cn)<2 or len(country)<2 or len(normalized)<8 or len(normalized)>15:c.close();flash(msg(lang,'profile_bad'));return redirect(url_for('account'))
                c.execute("UPDATE agencies SET contact_name=?,country=?,whatsapp=? WHERE id=?",(cn,country,wa,agency['id']));c.execute("UPDATE users SET name=?,mobile=? WHERE id=?",(cn,wa,u['id']));audit(c,u['id'],'agency_profile_update',f"agency={agency['id']}");c.commit();c.close();flash(msg(lang,'profile'));return redirect(url_for('account'))
            if action=='password':
                cp=request.form.get('current_password','');np=request.form.get('new_password','');confirm=request.form.get('confirm_password','')
                if not check_password_hash(u['password_hash'],cp):c.close();flash(msg(lang,'bad'));return redirect(url_for('account'))
                if len(np)<8:c.close();flash(msg(lang,'short'));return redirect(url_for('account'))
                if np!=confirm:c.close();flash(msg(lang,'mismatch'));return redirect(url_for('account'))
                if check_password_hash(u['password_hash'],np):c.close();flash(msg(lang,'same'));return redirect(url_for('account'))
                c.execute("UPDATE users SET password_hash=? WHERE id=?",(generate_password_hash(np),u['id']));audit(c,u['id'],'agency_password_change',f"agency={agency['id']}");c.commit();c.close();flash(msg(lang,'password'));return redirect(url_for('account'))
            c.close();return redirect(url_for('account'))
        if request.path=='/admin/settings' and request.method=='POST':
            u=current_user()
            if u and u['role']=='super_admin':
                c=db();values={'announcement_active':'1' if request.form.get('announcement_active') else '0','announcement_target':request.form.get('announcement_target','all') if request.form.get('announcement_target','all') in ('all','indonesia','malaysia') else 'all','announcement_start':request.form.get('announcement_start','').strip(),'announcement_end':request.form.get('announcement_end','').strip(),'announcement_whatsapp':'1' if request.form.get('announcement_whatsapp') else '0'}
                for key,value in values.items():c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(key,value))
                c.commit();c.close()
        u=current_user()
        if not u or u['role']!='staff':return None
        path=request.path;needed=None
        if path.startswith('/admin/reports'):needed='reports'
        elif path.startswith('/admin/agencies'):needed='agencies'
        elif path.startswith('/admin/hotels') or path.startswith('/admin/hotel/'):needed='hotels'
        elif path.startswith('/admin/offers'):needed='offers'
        if needed and not get_permissions(u['id']).get(needed,False):flash('ليس لديك صلاحية للوصول إلى هذا القسم.' if (u['language'] or 'ar')=='ar' else 'You do not have permission to access this section.');return redirect(url_for('admin'))
        return None

    @app.context_processor
    def extended_portal_context():
        if '_csrf_token' not in session:session['_csrf_token']=secrets.token_urlsafe(32)
        u=current_user();country='';staff_permissions={k:True for k in permission_keys};permissions={}
        if u and u['role']=='agency':
            c=db();row=c.execute("SELECT country FROM agencies WHERE user_id=?",(u['id'],)).fetchone();c.close();country=(row['country'] or '') if row else ''
        elif u and u['role']=='staff':staff_permissions=get_permissions(u['id'])
        elif u and u['role']=='super_admin':
            c=db();ensure_permissions_table(c);rows=c.execute("SELECT * FROM staff_permissions").fetchall();c.close();permissions={r['user_id']:{k:bool(r[k]) for k in permission_keys} for r in rows}
        return {'agency_country':country,'current_date':datetime.utcnow().date().isoformat(),'staff_permissions':staff_permissions,'permissions':permissions,'csrf_token':session['_csrf_token']}

    @app.route('/admin/employee/<int:uid>/permissions',methods=['POST'])
    def update_staff_permissions(uid):
        u=current_user()
        if not u or u['role']!='super_admin':return redirect(url_for('admin'))
        c=db();ensure_permissions_table(c);employee=c.execute("SELECT id,role FROM users WHERE id=?",(uid,)).fetchone()
        if not employee or employee['role']!='staff':c.close();return 'staff not found',404
        vals=[1 if request.form.get(k) else 0 for k in permission_keys];c.execute("INSERT INTO staff_permissions(user_id,reports,agencies,hotels,offers,updated_at) VALUES(?,?,?,?,?,?) ON CONFLICT(user_id) DO UPDATE SET reports=excluded.reports,agencies=excluded.agencies,hotels=excluded.hotels,offers=excluded.offers,updated_at=excluded.updated_at",(uid,*vals,datetime.utcnow().isoformat()));audit(c,u['id'],'staff_permissions_update',f"employee={uid}, reports={vals[0]}, agencies={vals[1]}, hotels={vals[2]}, offers={vals[3]}");c.commit();c.close();flash('تم تحديث صلاحيات الموظف.');return redirect(url_for('admin_employees'))

    @app.route('/admin/employee/<int:uid>/toggle',methods=['POST'])
    def toggle_staff(uid):
        u=current_user()
        if not u or u['role']!='super_admin':return redirect(url_for('admin'))
        c=db();row=c.execute("SELECT active,role FROM users WHERE id=?",(uid,)).fetchone()
        if not row or row['role']!='staff':c.close();return 'staff not found',404
        new_active=0 if row['active'] else 1;c.execute("UPDATE users SET active=? WHERE id=?",(new_active,uid));audit(c,u['id'],'staff_status_update',f"employee={uid}, active={new_active}");c.commit();c.close();flash('تم تحديث حالة الموظف.');return redirect(url_for('admin_employees'))

    def filters():return {'country':request.args.get('country','').strip(),'hotel':request.args.get('hotel','').strip(),'status':request.args.get('status','').strip(),'from':request.args.get('from','').strip(),'to':request.args.get('to','').strip()}
    def query_rows(c,f):
        sql="SELECT r.*,a.agency_name,a.country,a.contact_name,a.whatsapp,u.email,h.name_ar,h.name_en FROM requests r JOIN agencies a ON a.id=r.agency_id JOIN users u ON u.id=a.user_id LEFT JOIN hotels h ON h.id=r.hotel_id WHERE 1=1";params=[]
        if f['country']:sql+=' AND a.country LIKE ?';params.append('%'+f['country']+'%')
        if f['hotel']:sql+=" AND (h.name_ar LIKE ? OR h.name_en LIKE ? OR (r.any_hotel=1 AND ? LIKE '%any%'))";params+=['%'+f['hotel']+'%','%'+f['hotel']+'%',f['hotel'].lower()]
        if f['status'] in ('sent','contacted','closed'):sql+=' AND r.status=?';params.append(f['status'])
        if f['from']:sql+=' AND substr(r.created_at,1,10)>=?';params.append(f['from'])
        if f['to']:sql+=' AND substr(r.created_at,1,10)<=?';params.append(f['to'])
        return c.execute(sql+' ORDER BY r.id DESC',params).fetchall()

    @app.route('/admin/reports')
    @admin_required
    def admin_reports():
        f=filters();c=db();rows=query_rows(c,f);summary={'requests':len(rows),'rooms':sum((r['rooms'] or 0) for r in rows),'persons':sum((r['persons'] or 0) for r in rows),'agencies':len(set(r['agency_id'] for r in rows))};countries=c.execute("SELECT country,COUNT(*) total FROM agencies GROUP BY country ORDER BY total DESC,country").fetchall();hotels=c.execute("SELECT COALESCE(h.name_en,'Any Available Hotel') hotel,COUNT(*) total FROM requests r LEFT JOIN hotels h ON h.id=r.hotel_id GROUP BY COALESCE(h.name_en,'Any Available Hotel') ORDER BY total DESC").fetchall();c.close();return render_template('admin_reports.html',user=current_user(),rows=rows,summary=summary,filters=f,countries=countries,hotels=hotels)

    @app.route('/admin/reports.xlsx')
    @admin_required
    def admin_reports_excel():
        f=filters();c=db();rows=query_rows(c,f);c.close();wb=Workbook();ws=wb.active;ws.title='Requests';headers=['Request ID','Agency','Country','Contact','WhatsApp','Email','Hotel','City','Check-in','Check-out','Rooms','Persons','Nationality','Meal','Status','Created At'];ws.append(headers)
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='111111');cell.alignment=Alignment(horizontal='center')
        for r in rows:ws.append([r['id'],r['agency_name'],r['country'],r['contact_name'],r['whatsapp'],r['email'],'Any Available Hotel' if r['any_hotel'] else (r['name_en'] or r['name_ar'] or ''),r['city'],r['checkin'],r['checkout'],r['rooms'],r['persons'],r['nationality'],r['meal'],r['status'],r['created_at']])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for i in range(1,len(headers)+1):max_len=max(len(str(ws.cell(row=j,column=i).value or '')) for j in range(1,min(ws.max_row,500)+1));ws.column_dimensions[get_column_letter(i)].width=min(max(max_len+2,12),32)
        summary=wb.create_sheet('Summary');summary.append(['MUROOJ GOLDEN B2B','Report Summary']);summary.append(['Requests',len(rows)]);summary.append(['Rooms',sum((r['rooms'] or 0) for r in rows)]);summary.append(['Persons',sum((r['persons'] or 0) for r in rows)]);summary.append(['Agencies',len(set(r['agency_id'] for r in rows))]);summary['A1'].font=Font(bold=True);summary['B1'].font=Font(bold=True);out=BytesIO();wb.save(out);out.seek(0);return send_file(out,as_attachment=True,download_name='murooj-golden-full-report.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    from offer_extensions import register_offer_extensions
    register_offer_extensions(app,db,current_user)
    from device_security import register_device_security
    register_device_security(app,db,current_user)
